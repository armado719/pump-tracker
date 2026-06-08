from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Estilos globales ──────────────────────────────────────────────────────────
style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(11)

def set_heading(paragraph, level, text, color=None):
    paragraph.text = text
    paragraph.style = doc.styles[f'Heading {level}']
    if color:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(*color)

def add_heading(doc, level, text, color=None):
    p = doc.add_heading(text, level=level)
    if color:
        for run in p.runs:
            run.font.color.rgb = RGBColor(*color)
    return p

def add_image_placeholder(doc, number, title, description):
    """Agrega un bloque visual de espacio para imagen."""
    doc.add_paragraph()
    # Borde/caja simulada con tabla de 1 celda
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    cell.width = Inches(6)

    # Título de la imagen
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_title.add_run(f'[ IMAGEN {number} — {title} ]')
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    # Espacio visual para la imagen
    p_space = cell.add_paragraph()
    p_space.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_space = p_space.add_run('\n\n\n\n\n\n')
    run_space.font.size = Pt(8)

    p_desc = cell.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_desc = p_desc.add_run(f'↑  {description}  ↑')
    run_desc.font.italic = True
    run_desc.font.size = Pt(9)
    run_desc.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    doc.add_paragraph()

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    # Encabezados
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Fondo oscuro para encabezado
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1F497D')
        tcPr.append(shd)
    # Filas
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, cell_text in enumerate(row_data):
            row.cells[c_idx].text = str(cell_text)
            if (r_idx % 2) == 0:
                tc = row.cells[c_idx]._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'DCE6F1')
                tcPr.append(shd)
    doc.add_paragraph()
    return table

BLUE = (0x1F, 0x49, 0x7D)
DARK = (0x26, 0x26, 0x26)

# ═══════════════════════════════════════════════════════════════════════════════
# PORTADA
# ═══════════════════════════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

p_logo = doc.add_paragraph()
p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p_logo.add_run('GRS')
run.font.size = Pt(48)
run.font.bold = True
run.font.color.rgb = RGBColor(*BLUE)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = p_sub.add_run('General Rigs Services S.A.S.')
run2.font.size = Pt(16)
run2.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

doc.add_paragraph()

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run3 = p_title.add_run('PUMP TRACKER GRS')
run3.font.size = Pt(32)
run3.font.bold = True
run3.font.color.rgb = RGBColor(*BLUE)

p_subtitle = doc.add_paragraph()
p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run4 = p_subtitle.add_run('Manual de Usuario')
run4.font.size = Pt(20)
run4.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

p_ver = doc.add_paragraph()
p_ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
run5 = p_ver.add_run('Versión 1.0  —  Junio 2026')
run5.font.size = Pt(12)
run5.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 1. INTRODUCCIÓN
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '1. Introducción', BLUE)
doc.add_paragraph(
    'Pump Tracker GRS es una aplicación web de seguimiento operacional desarrollada para '
    'General Rigs Services S.A.S. Permite registrar y monitorear en tiempo real el estado '
    'de los equipos de perforación (rigs), las bombas de lodo, los componentes críticos, '
    'el cable de perforación TM y las horas de operación diarias.'
)

add_heading(doc, 2, '1.1 Roles del Sistema', BLUE)
add_table(doc,
    ['Rol', 'Permisos'],
    [
        ['Administrador', 'Acceso total: crear, editar, eliminar registros y gestionar usuarios.'],
        ['Operador', 'Consulta y registro de operaciones diarias. No puede crear ni eliminar equipos.'],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 2. ACCESO AL SISTEMA
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '2. Acceso al Sistema', BLUE)
add_heading(doc, 2, '2.1 Pantalla de Login', BLUE)
doc.add_paragraph(
    'La pantalla de inicio de sesión es el punto de entrada al sistema. '
    'El usuario debe ingresar su correo electrónico y contraseña registrados.'
)

add_image_placeholder(doc, 1, 'Pantalla de Login',
    'Captura de la pantalla de inicio de sesión con logo GRS, campos de correo, contraseña y botón Iniciar sesión')

add_heading(doc, 3, 'Campos', BLUE)
add_table(doc,
    ['Campo', 'Descripción'],
    [
        ['Correo electrónico', 'Dirección de correo asociada a la cuenta.'],
        ['Contraseña', 'Contraseña asignada por el administrador.'],
        ['Recordarme', 'Mantiene la sesión activa en el navegador.'],
    ]
)
doc.add_paragraph(
    'Nota: Si olvidó su contraseña use el enlace "¿Olvidaste tu contraseña?" '
    'para recibir un correo de recuperación.'
).runs[0].font.italic = True

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 3. DASHBOARD PRINCIPAL
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '3. Dashboard Principal', BLUE)
doc.add_paragraph(
    'El Dashboard es la pantalla principal del sistema. Muestra un resumen general '
    'del estado operacional de todos los equipos en tiempo real.'
)

add_image_placeholder(doc, 2, 'Dashboard Principal',
    'Captura completa del dashboard: tarjetas de resumen, tabla Estado de Rigs, Alertas Recientes y gráfico de Horas de Trabajo')

add_heading(doc, 2, '3.1 Tarjetas de Resumen', BLUE)
add_table(doc,
    ['Tarjeta', 'Descripción'],
    [
        ['Rigs Activos', 'Número total de equipos registrados en el sistema.'],
        ['Bombas Operando Hoy', 'Bombas con registro de horas en el día actual.'],
        ['Alertas Críticas', 'Componentes en estado crítico que requieren atención inmediata.'],
        ['Horas Prom. (7D)', 'Promedio de horas trabajadas por las bombas en los últimos 7 días.'],
        ['Cables TM — Alertas', 'Cables de perforación que superaron el umbral de alerta configurado.'],
    ]
)

add_heading(doc, 2, '3.2 Estado de Rigs', BLUE)
doc.add_paragraph(
    'Tabla que muestra por cada equipo: el pozo asignado, número de bombas, horas trabajadas hoy, '
    'fecha de última actualización y estado (Activo / Inactivo).'
)

add_heading(doc, 2, '3.3 Gráfico de Horas de Trabajo', BLUE)
doc.add_paragraph(
    'Gráfico de líneas con las horas diarias trabajadas por los rigs en los últimos 30 días.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 4. GESTIÓN DE EQUIPOS (RIGS)
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '4. Gestión de Equipos (Rigs)', BLUE)

add_heading(doc, 2, '4.1 Lista de Equipos', BLUE)
add_image_placeholder(doc, 3, 'Lista de Rigs',
    'Captura de la tabla de equipos mostrando nombre del rig y acciones disponibles')

add_heading(doc, 2, '4.2 Crear Equipo  (solo Administrador)', BLUE)
doc.add_paragraph('Formulario para registrar un nuevo equipo de perforación en el sistema.')
add_image_placeholder(doc, 4, 'Formulario Crear Rig',
    'Captura del formulario de creación de rig con el campo Nombre del equipo')

add_heading(doc, 2, '4.3 Detalle del Equipo', BLUE)
doc.add_paragraph(
    'Vista detallada de un equipo específico. Muestra las bombas asociadas, '
    'los pozos asignados y el historial de operaciones.'
)
add_image_placeholder(doc, 5, 'Detalle del Rig',
    'Captura de la página de detalle de un rig mostrando sus bombas asociadas y pozos')

add_heading(doc, 2, '4.4 Editar Equipo  (solo Administrador)', BLUE)
add_image_placeholder(doc, 6, 'Editar Rig',
    'Captura del formulario de edición de rig')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 5. GESTIÓN DE BOMBAS
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '5. Gestión de Bombas', BLUE)

add_heading(doc, 2, '5.1 Lista de Bombas', BLUE)
add_image_placeholder(doc, 7, 'Lista de Bombas',
    'Captura de la tabla de bombas con columnas de rig, número de bomba y estado general')

add_heading(doc, 2, '5.2 Crear Bomba  (solo Administrador)', BLUE)
add_image_placeholder(doc, 8, 'Formulario Crear Bomba',
    'Captura del formulario de creación de bomba mostrando campos de equipo, número y tipo')

add_heading(doc, 2, '5.3 Detalle de Bomba', BLUE)
doc.add_paragraph(
    'Vista completa de una bomba. Muestra los ensambles (izquierdo, derecho, superior) '
    'con sus componentes, las horas acumuladas de cada pieza y el estado.'
)
add_image_placeholder(doc, 9, 'Detalle de Bomba con Componentes',
    'Captura mostrando los tres ensambles con componentes, barras de progreso de horas y badges de estado')

add_heading(doc, 3, 'Estados de Componentes', BLUE)
add_table(doc,
    ['Estado', 'Color', 'Descripción'],
    [
        ['OK', 'Verde', 'Horas acumuladas dentro del rango seguro.'],
        ['Alerta', 'Amarillo/Naranja', 'Próximo al límite recomendado.'],
        ['Crítico', 'Rojo', 'Superó el límite; requiere reemplazo inmediato.'],
    ]
)

add_heading(doc, 2, '5.4 Historial de Reemplazos', BLUE)
add_image_placeholder(doc, 10, 'Historial de Reemplazos',
    'Captura del historial mostrando fecha, componente reemplazado y horas al momento del reemplazo')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 6. REGISTRO DIARIO DE OPERACIONES
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '6. Registro Diario de Operaciones', BLUE)

add_heading(doc, 2, '6.1 Lista de Registros', BLUE)
add_image_placeholder(doc, 11, 'Lista de Registros Diarios',
    'Captura de la tabla de registros diarios con columnas de fecha, pozo, horas trabajadas y observaciones')

add_heading(doc, 2, '6.2 Crear Registro Diario', BLUE)
add_image_placeholder(doc, 12, 'Formulario Registro Diario',
    'Captura del formulario mostrando campos de fecha, pozo, horas trabajadas y observaciones')

add_heading(doc, 3, 'Campos', BLUE)
add_table(doc,
    ['Campo', 'Descripción'],
    [
        ['Fecha', 'Día del registro (por defecto: hoy).'],
        ['Pozo', 'Pozo en el que operó la bomba ese día.'],
        ['Horas trabajadas', 'Horas de operación efectiva.'],
        ['Observaciones', 'Notas adicionales del operador.'],
    ]
)

add_heading(doc, 2, '6.3 Detalle del Registro', BLUE)
add_image_placeholder(doc, 13, 'Detalle del Registro Diario',
    'Captura de la vista de detalle de un registro diario')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 7. REEMPLAZO DE COMPONENTES
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '7. Reemplazo de Componentes', BLUE)
doc.add_paragraph(
    'Cuando un componente llega a estado crítico, el operador puede registrar su '
    'reemplazo desde la página de detalle de la bomba.'
)
add_image_placeholder(doc, 14, 'Formulario de Reemplazo de Componente',
    'Captura del formulario de reemplazo mostrando el componente, horas actuales y botón de confirmación')

p = doc.add_paragraph('Al registrar un reemplazo:')
doc.add_paragraph('• Las horas acumuladas del componente se reinician a cero.', style='List Bullet')
doc.add_paragraph(
    '• Queda registrado en el historial de reemplazos con la fecha y las horas al momento del cambio.',
    style='List Bullet'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 8. CENTRO DE ALERTAS
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '8. Centro de Alertas', BLUE)
doc.add_paragraph(
    'Pantalla centralizada que muestra todos los componentes y cables TM que requieren atención.'
)
add_image_placeholder(doc, 15, 'Centro de Alertas',
    'Captura del centro de alertas con tarjetas de resumen y tablas de alertas de bombas y cables TM')

add_heading(doc, 2, '8.1 Secciones', BLUE)
add_table(doc,
    ['Sección', 'Descripción'],
    [
        ['Alertas de Bombas — Componentes', 'Lista de componentes en estado alerta o crítico con rig, bomba, ensamble y horas.'],
        ['Alertas de Cable TM', 'Cables que superaron el porcentaje de alerta configurado.'],
    ]
)

add_heading(doc, 2, '8.2 Notificación por Correo', BLUE)
doc.add_paragraph(
    'El botón "Notificar por email" envía un resumen de todas las alertas activas '
    'al correo configurado en el sistema.'
)
add_image_placeholder(doc, 16, 'Botón Notificar y Confirmación de Envío',
    'Captura del botón de notificación y mensaje de confirmación tras el envío exitoso')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 9. REPORTES PDF
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '9. Reportes PDF', BLUE)
add_image_placeholder(doc, 17, 'Panel de Reportes',
    'Captura del panel de reportes con opciones de reporte de bombas y reporte de cable TM')

add_table(doc,
    ['Reporte', 'Contenido'],
    [
        ['Reporte de Bombas', 'Estado de todos los componentes de todas las bombas de un rig, con horas y estado.'],
        ['Reporte Cable TM', 'Resumen de TM acumuladas, operaciones registradas e historial del cable activo.'],
    ]
)

add_image_placeholder(doc, 18, 'Ejemplo de Reporte PDF Generado',
    'Captura del PDF generado abierto en el navegador, mostrando portada y tabla de componentes')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 10. CABLE TM
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '10. Cable TM', BLUE)
doc.add_paragraph(
    'Módulo especializado para el seguimiento del cable de perforación mediante el '
    'cálculo de Ton-Millas (TM) acumuladas.'
)

add_heading(doc, 2, '10.1 Dashboard Cable TM', BLUE)
add_image_placeholder(doc, 19, 'Dashboard Cable TM — Sin cable activo',
    'Captura del dashboard cuando no hay cable registrado, mostrando el formulario de registro')

add_image_placeholder(doc, 20, 'Dashboard Cable TM — Con cable activo',
    'Captura con cable activo: barra de progreso TM, datos del cable y tabla de últimas operaciones')

add_heading(doc, 2, '10.2 Nueva Operación TM', BLUE)
add_image_placeholder(doc, 21, 'Formulario Nueva Operación TM',
    'Captura del formulario de nueva operación con campos de tipo, peso bloque, líneas, profundidad y distancia')

add_heading(doc, 3, 'Campos', BLUE)
add_table(doc,
    ['Campo', 'Descripción'],
    [
        ['Tipo de operación', 'Tipo de maniobra realizada (perforación, viaje, etc.)'],
        ['Peso del bloque (lb)', 'Carga sobre el gancho.'],
        ['N° líneas activas', 'Líneas en el sistema de poleas.'],
        ['Profundidad (ft)', 'Profundidad del pozo.'],
        ['Distancia recorrida (ft)', 'Longitud de cable movida en la operación.'],
    ]
)

add_heading(doc, 2, '10.3 Historial de Cables', BLUE)
add_image_placeholder(doc, 22, 'Historial de Cables',
    'Captura del historial mostrando tabla de cables con fechas de instalación, TM acumuladas y estado')

add_heading(doc, 2, '10.4 Configuración del Equipo', BLUE)
add_image_placeholder(doc, 23, 'Formulario de Configuración Cable TM',
    'Captura del formulario con todos los parámetros: altura torre, diámetro tambor, líneas, TM máx., umbral')

add_heading(doc, 3, 'Parámetros Configurables', BLUE)
add_table(doc,
    ['Campo', 'Descripción', 'Default'],
    [
        ['Altura Torre (ft)', 'Altura de la torre de perforación.', '105 ft'],
        ['Diámetro Tambor', 'Diámetro del tambor del malacate (in / cm / ft).', '18 in'],
        ['N° Líneas Activas', 'Número de líneas en el sistema de poleas.', '8'],
        ['TM Máx. para Corte', 'Ton-millas máximas antes de reemplazar el cable.', '1200'],
        ['Tipo de Cable', 'Clasificación del cable (EIP / EEIP / IPS).', 'EIP'],
        ['Peso Bloque Default (lb)', 'Peso del bloque de viaje por defecto.', '25,000 lb'],
        ['Umbral de Alerta (%)', 'Porcentaje del TM máximo que activa la alerta.', '80%'],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 11. PANEL GERENCIAL
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '11. Panel Gerencial', BLUE)
doc.add_paragraph(
    'Vista ejecutiva con indicadores clave de desempeño (KPIs) consolidados de todos los equipos. '
    'Diseñada para gerencia y supervisión de alto nivel.'
)
add_image_placeholder(doc, 24, 'Panel Gerencial',
    'Captura del panel gerencial mostrando gráficas de producción y métricas de eficiencia operacional')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 12. ADMINISTRACIÓN
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '12. Administración', BLUE)
p = doc.add_paragraph()
run = p.add_run('Las siguientes secciones son exclusivas para usuarios con rol Administrador.')
run.font.bold = True
run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

add_heading(doc, 2, '12.1 Gestión de Usuarios', BLUE)
add_image_placeholder(doc, 25, 'Lista de Usuarios',
    'Captura de la tabla de usuarios con nombre, correo, rol, estado y botones de acción')

add_image_placeholder(doc, 26, 'Formulario Crear/Editar Usuario',
    'Captura del formulario de usuario con campos de nombre, correo, rol y contraseña')

add_heading(doc, 3, 'Acciones Disponibles', BLUE)
for accion in [
    'Crear usuario: registrar nuevo acceso al sistema.',
    'Editar: modificar nombre, correo o rol.',
    'Activar / Desactivar: controlar el acceso sin eliminar el historial.',
    'Eliminar: remover permanentemente el usuario.',
]:
    doc.add_paragraph(accion, style='List Bullet')

add_heading(doc, 2, '12.2 Gestión de Pozos', BLUE)
add_image_placeholder(doc, 27, 'Lista de Pozos',
    'Captura de la tabla de pozos con nombre y acciones de editar/eliminar')

add_image_placeholder(doc, 28, 'Formulario Crear/Editar Pozo',
    'Captura del formulario de pozo mostrando el campo de nombre')

add_heading(doc, 2, '12.3 Configuración de Correo SMTP', BLUE)
doc.add_paragraph(
    'Configura el servidor de correo que usa el sistema para enviar notificaciones de alertas.'
)
add_image_placeholder(doc, 29, 'Configuración SMTP',
    'Captura del formulario SMTP con campos de servidor, puerto, usuario y encriptación')

add_heading(doc, 3, 'Campos', BLUE)
add_table(doc,
    ['Campo', 'Descripción'],
    [
        ['Host SMTP', 'Servidor de correo saliente (ej. smtp.gmail.com).'],
        ['Puerto', 'Puerto del servidor (ej. 587 para TLS).'],
        ['Usuario', 'Dirección de correo remitente.'],
        ['Encriptación', 'TLS o SSL.'],
        ['Nombre remitente', 'Nombre que aparece en los correos enviados.'],
        ['Correo destino de alertas', 'Dirección que recibe las notificaciones.'],
    ]
)

add_heading(doc, 2, '12.4 Auditoría del Sistema', BLUE)
doc.add_paragraph(
    'Registro detallado de todas las acciones realizadas por los usuarios: '
    'creaciones, ediciones y eliminaciones de registros.'
)
add_image_placeholder(doc, 30, 'Log de Auditoría',
    'Captura del log de auditoría con columnas de fecha, usuario, acción y modelo afectado')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 13. PERFIL Y CONTRASEÑA
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, '13. Perfil y Contraseña', BLUE)
doc.add_paragraph(
    'Cada usuario puede actualizar su información de perfil y cambiar su contraseña '
    'desde el menú inferior del sidebar.'
)
add_image_placeholder(doc, 31, 'Perfil de Usuario',
    'Captura de la página de perfil con formulario de actualización de nombre, correo y cambio de contraseña')

add_table(doc,
    ['Opción', 'Descripción'],
    [
        ['Actualizar información', 'Cambiar nombre y dirección de correo.'],
        ['Cambiar contraseña', 'Requiere ingresar la contraseña actual para confirmar.'],
        ['Eliminar cuenta', 'Opción permanente que requiere confirmación con contraseña.'],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# GLOSARIO
# ═══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 1, 'Glosario', BLUE)
add_table(doc,
    ['Término', 'Definición'],
    [
        ['Rig', 'Equipo de perforación terrestre.'],
        ['Bomba de lodo', 'Bomba de alta presión usada en la circulación del fluido de perforación.'],
        ['Ensamble', 'Conjunto de componentes de una bomba (izquierdo, derecho, superior).'],
        ['TM / Ton-Milla', 'Unidad de medida de desgaste del cable de perforación.'],
        ['Umbral', 'Límite de horas/TM a partir del cual se genera una alerta.'],
        ['Malacate', 'Equipo de izaje que enrolla el cable de perforación.'],
        ['EIP / EEIP / IPS', 'Clasificaciones de resistencia del cable de perforación.'],
    ]
)

# Pie de página
doc.add_paragraph()
p_footer = doc.add_paragraph()
p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_f = p_footer.add_run(
    'Documento para uso interno de General Rigs Services S.A.S.\n'
    'Pump Tracker GRS — Todos los derechos reservados © 2026'
)
run_f.font.italic = True
run_f.font.size = Pt(9)
run_f.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

# Guardar
output = '/home/user/pump-tracker/docs/Manual_Usuario_PumpTracker_GRS.docx'
doc.save(output)
print(f'Documento generado: {output}')
